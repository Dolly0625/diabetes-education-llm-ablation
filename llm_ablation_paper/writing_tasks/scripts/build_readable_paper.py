#!/usr/bin/env python3
"""Build a readable DOCX edition of the Chinese conference-paper draft."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "drafts" / "PAPER_FULL_DRAFT_ZH.md"
BIB = ROOT / "drafts" / "references.bib"
FIGURES = ROOT / "drafts" / "figures"

BODY_FONT_CJK = "Noto Serif CJK TC"
BODY_FONT_LATIN = "Times New Roman"
SANS_FONT_CJK = "Noto Sans CJK TC"
SANS_FONT_LATIN = "Arial"
MONO_FONT = "Menlo"


def set_run_font(run, latin=BODY_FONT_LATIN, cjk=BODY_FONT_CJK, size=None, bold=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    set_run_font(run, size=9)


def parse_bib(path: Path):
    text = path.read_text(encoding="utf-8")
    entries = {}
    starts = list(re.finditer(r"@(\w+)\s*\{\s*([^,]+),", text))
    for idx, match in enumerate(starts):
        start = match.end()
        end = starts[idx + 1].start() if idx + 1 < len(starts) else len(text)
        body = text[start:end].rstrip().rstrip("}").strip()
        fields = {}
        for fm in re.finditer(r"(\w+)\s*=\s*(?:\{((?:[^{}]|\{[^{}]*\})*)\}|\"([^\"]*)\")\s*,?", body, re.S):
            value = fm.group(2) if fm.group(2) is not None else fm.group(3)
            fields[fm.group(1).lower()] = re.sub(r"\s+", " ", value).strip()
        entries[match.group(2).strip()] = fields
    return entries


def citation_order(text: str):
    order = []
    seen = set()
    for group in re.findall(r"\[@([^\]]+)\]", text):
        for raw in group.split(";"):
            key = raw.strip().lstrip("@")
            if key and key not in seen:
                seen.add(key)
                order.append(key)
    return order


def clean_bib_value(value: str):
    value = value.replace("{", "").replace("}", "")
    value = value.replace("--", "-")
    return value


def format_reference(fields: dict):
    authors = clean_bib_value(fields.get("author", ""))
    authors = authors.replace(" and ", "; ")
    year = fields.get("year", "n.d.")
    title = clean_bib_value(fields.get("title", "Untitled"))
    container = clean_bib_value(fields.get("journal") or fields.get("booktitle") or "")
    volume = fields.get("volume", "")
    number = fields.get("number", "")
    pages = fields.get("pages", "")
    doi = fields.get("doi", "")
    url = fields.get("url", "")
    parts = [f"{authors} ({year}). {title}."]
    if container:
        tail = container
        if volume:
            tail += f", {volume}"
        if number:
            tail += f"({number})"
        if pages:
            tail += f", {pages}"
        parts.append(tail + ".")
    if doi:
        parts.append(f"https://doi.org/{doi}")
    elif url:
        parts.append(url)
    return " ".join(parts)


def replace_citations(text: str, citation_numbers: dict):
    def repl(match):
        nums = []
        for raw in match.group(1).split(";"):
            key = raw.strip().lstrip("@")
            if key in citation_numbers:
                nums.append(str(citation_numbers[key]))
        return "[" + ", ".join(nums) + "]" if nums else ""

    return re.sub(r"\[@([^\]]+)\]", repl, text)


INLINE_RE = re.compile(r"(\*\*.*?\*\*|`.*?`|\*[^*]+?\*)")


def add_inline_runs(paragraph, text: str, citation_numbers: dict):
    text = replace_citations(text, citation_numbers)
    text = text.replace("{FIG:", "圖 ").replace("{TAB:", "表 ").replace("}", "")
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=10.5)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=10.5, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, latin=MONO_FONT, cjk=MONO_FONT, size=9)
            run.font.color.rgb = RGBColor(60, 60, 60)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=10.5)
            run.italic = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=10.5)


def style_paragraph(paragraph, first_line=True, after=5, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing = line
    fmt.space_after = Pt(after)
    if first_line:
        fmt.first_line_indent = Cm(0.74)
    fmt.widow_control = True


def add_body_paragraph(doc, text, citation_numbers, first_line=True, style=None):
    p = doc.add_paragraph(style=style)
    add_inline_runs(p, text, citation_numbers)
    style_paragraph(p, first_line=first_line)
    return p


def add_caption(doc, text, citation_numbers):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    add_inline_runs(p, text, citation_numbers)
    for run in p.runs:
        set_run_font(run, size=9)
    return p


def add_figure(doc, image_path: Path, max_width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(max_width))
    p.paragraph_format.keep_with_next = True


def add_table(doc, rows, citation_numbers, widths=None):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 or len(row) <= 3 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline_runs(p, value, citation_numbers)
            for run in p.runs:
                set_run_font(run, latin=BODY_FONT_LATIN, cjk=BODY_FONT_CJK, size=8.2, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "24364B")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F2F5F8")
            if widths and c_idx < len(widths):
                cell.width = Inches(widths[c_idx])
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_CJK)

    title = styles["Title"]
    title.font.name = SANS_FONT_LATIN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), SANS_FONT_CJK)
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 15, 14, 7),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = SANS_FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), SANS_FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_front_matter(doc):
    p = doc.add_paragraph(style="Title")
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(85)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("糖尿病衛教大型語言模型之分層防護網消融研究")
    set_run_font(r, latin=SANS_FONT_LATIN, cjk=SANS_FONT_CJK, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("以對抗性壓力測試評估系統安全與錯誤再分配")
    set_run_font(r, latin=SANS_FONT_LATIN, cjk=SANS_FONT_CJK, size=13, bold=False)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("研討會論文中文閱讀版")
    set_run_font(r, latin=SANS_FONT_LATIN, cjk=SANS_FONT_CJK, size=10)
    r.font.color.rgb = RGBColor(90, 90, 90)
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build(output: Path):
    source_text = SOURCE.read_text(encoding="utf-8")
    bib_entries = parse_bib(BIB)
    citation_keys = citation_order(source_text)
    citation_numbers = {key: i + 1 for i, key in enumerate(citation_keys)}
    lines = source_text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.85)
    configure_styles(doc)
    add_front_matter(doc)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("糖尿病衛教 LLM 分層防護消融研究")
    set_run_font(hr, latin=SANS_FONT_LATIN, cjk=SANS_FONT_CJK, size=8)
    hr.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])

    i = 1  # source title is rendered as front matter
    in_mermaid = False
    mermaid_inserted = False
    references_written = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line == "```mermaid":
            in_mermaid = True
            i += 1
            continue
        if in_mermaid:
            if line == "```":
                in_mermaid = False
                add_figure(doc, FIGURES / "figure1_system_ablation.png", max_width=6.2)
                mermaid_inserted = True
            i += 1
            continue
        if line in ("", "---"):
            i += 1
            continue
        if line.startswith("<!--"):
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if line.startswith("::: {#refs}") or line == ":::":
            i += 1
            continue
        if line.startswith("# "):
            heading = line[2:].strip()
            if heading.startswith("參考文獻"):
                doc.add_page_break()
                p = doc.add_paragraph("參考文獻", style="Heading 1")
                p.paragraph_format.keep_with_next = True
                for key in citation_keys:
                    fields = bib_entries.get(key, {})
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.75)
                    p.paragraph_format.first_line_indent = Cm(-0.75)
                    p.paragraph_format.space_after = Pt(4)
                    r = p.add_run(f"[{citation_numbers[key]}] ")
                    set_run_font(r, size=9)
                    r = p.add_run(format_reference(fields))
                    set_run_font(r, size=9)
                references_written = True
            else:
                doc.add_paragraph(heading, style="Heading 1")
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 3")
            i += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            add_figure(doc, ROOT / "drafts" / image_match.group(2), max_width=6.25)
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            widths = None
            if rows:
                if len(rows[0]) == 5:
                    widths = [1.05, 0.75, 1.15, 1.05, 3.05]
                elif len(rows[0]) == 3:
                    widths = [1.45, 4.65, 0.65]
            add_table(doc, rows, citation_numbers, widths)
            continue
        if re.match(r"^\*\*\{(?:FIG|TAB):\d+\}", line):
            add_caption(doc, line, citation_numbers)
            i += 1
            continue
        if line.startswith("*表註："):
            p = add_body_paragraph(doc, line.strip("*"), citation_numbers, first_line=False)
            for run in p.runs:
                set_run_font(run, size=8.5)
                run.italic = True
                run.font.color.rgb = RGBColor(70, 70, 70)
            i += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)", raw)
        if bullet:
            p = add_body_paragraph(doc, bullet.group(1), citation_numbers, first_line=False, style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.25)
            i += 1
            continue
        numbered = re.match(r"^\s*(\d+)\.\s+(.*)", raw)
        if numbered:
            p = doc.add_paragraph()
            prefix = p.add_run(f"{numbered.group(1)}. ")
            set_run_font(prefix, size=10.5)
            add_inline_runs(p, numbered.group(2), citation_numbers)
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.25)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            i += 1
            continue
        add_body_paragraph(doc, line, citation_numbers, first_line=True)
        i += 1

    if not mermaid_inserted:
        raise RuntimeError("Figure 1 was not inserted")
    if not references_written:
        raise RuntimeError("References were not generated")

    core = doc.core_properties
    core.title = "糖尿病衛教大型語言模型之分層防護網消融研究"
    core.subject = "研討會論文中文閱讀版"
    core.author = ""
    core.keywords = "糖尿病衛教, 大型語言模型, 消融研究, 安全評估"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: build_readable_paper.py OUTPUT.docx")
    build(Path(sys.argv[1]).resolve())
